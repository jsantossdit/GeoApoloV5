"""
Gerador de Arquivos de Relatórios (Excel, HTML/PDF e CSV).
"""

import csv
import os
import tempfile
import webbrowser
from datetime import datetime
from typing import List, Dict, Any
from relatorios.models import TipoRelatorio

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class RelatorioGenerator:

    @staticmethod
    def gerar_excel(dados: List[Dict[str, Any]], titulo: str, caminho_destino: str) -> str:
        """Gera planilha Excel profissional com estilos corporativos usando openpyxl."""
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("Biblioteca openpyxl não está instalada no ambiente.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatório"

        if not dados:
            ws["A1"] = "Nenhum registro encontrado para os filtros selecionados."
            wb.save(caminho_destino)
            return caminho_destino

        colunas = list(dados[0].keys())

        # Estilos
        cor_cabecalho = "1A365D"  # Azul Marinho Corporativo
        fill_cabecalho = PatternFill(start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid")
        font_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_titulo = Font(name="Segoe UI", size=14, bold=True, color="1A365D")
        font_dados = Font(name="Segoe UI", size=10)
        border_fina = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )
        fill_zebra = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")

        # Título do Relatório
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colunas))
        ws["A1"] = f"GEOALVO - {titulo.upper()}"
        ws["A1"].font = font_titulo
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 30

        # Subtítulo com Data/Hora de Geração
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(colunas))
        ws["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')} | Total de Registros: {len(dados)}"
        ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="718096")
        ws.row_dimensions[2].height = 18

        # Cabeçalhos das Colunas (Linha 4)
        linha_cabecalho = 4
        ws.row_dimensions[linha_cabecalho].height = 24
        for col_idx, col_nome in enumerate(colunas, 1):
            cell = ws.cell(row=linha_cabecalho, column=col_idx, value=col_nome)
            cell.font = font_cabecalho
            cell.fill = fill_cabecalho
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_fina

        # Linhas de Dados
        for row_idx, reg in enumerate(dados, start=linha_cabecalho + 1):
            ws.row_dimensions[row_idx].height = 20
            usar_zebra = (row_idx % 2 == 0)
            for col_idx, col_nome in enumerate(colunas, 1):
                valor = reg.get(col_nome, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=valor)
                cell.font = font_dados
                cell.border = border_fina
                if usar_zebra:
                    cell.fill = fill_zebra

        # Ajuste automático das larguras de coluna
        for col_idx, col_nome in enumerate(colunas, 1):
            max_len = max(
                len(str(col_nome)),
                max((len(str(r.get(col_nome, ""))) for r in dados), default=0)
            )
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        wb.save(caminho_destino)
        return caminho_destino

    @staticmethod
    def gerar_html_imprimivel(dados: List[Dict[str, Any]], titulo: str, abrir_navegador: bool = True) -> str:
        """Gera relatório HTML pronto para impressão / PDF e abre no navegador padrão."""
        colunas = list(dados[0].keys()) if dados else []

        linhas_html = []
        for r in dados:
            tds = "".join(f"<td>{str(r.get(c, '') or '')}</td>" for c in colunas)
            linhas_html.append(f"<tr>{tds}</tr>")

        cabecalhos_html = "".join(f"<th>{c}</th>" for c in colunas)

        html_content = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
    <meta charset="UTF-8">
    <title>GeoAlvo - {titulo}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; margin: 20px; color: #2D3748; }}
        .header {{ border-bottom: 2px solid #1A365D; padding-bottom: 10px; margin-bottom: 20px; }}
        h1 {{ margin: 0 0 5px 0; font-size: 20px; color: #1A365D; }}
        .meta {{ font-size: 12px; color: #718096; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 11px; }}
        th {{ background-color: #1A365D; color: white; padding: 8px; text-align: left; border: 1px solid #CBD5E0; }}
        td {{ padding: 6px 8px; border: 1px solid #E2E8F0; }}
        tr:nth-child(even) {{ background-color: #F7FAFC; }}
        @media print {{
            button {{ display: none; }}
            body {{ margin: 0; }}
        }}
        .btn-print {{ background: #2B6CB0; color: white; border: none; padding: 8px 16px; border-radius: 4px; cursor: pointer; margin-bottom: 15px; font-size: 13px; }}
    </style>
</head>
<body>
    <button class="btn-print" onclick="window.print()">🖨 Imprimir / Salvar como PDF</button>
    <div class="header">
        <h1>GEOALVO - {titulo.upper()}</h1>
        <div class="meta">Gerado em: {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')} | Total: {len(dados)} registro(s)</div>
    </div>
    <table>
        <thead>
            <tr>{cabecalhos_html}</tr>
        </thead>
        <tbody>
            {"".join(linhas_html)}
        </tbody>
    </table>
</body>
</html>
"""
        caminho = os.path.join(tempfile.gettempdir(), f"relatorio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html")
        with open(caminho, "w", encoding="utf-8") as f:
            f.write(html_content)

        if abrir_navegador:
            webbrowser.open(caminho)
        return caminho

    @staticmethod
    def gerar_csv(dados: List[Dict[str, Any]], caminho_destino: str) -> str:
        """Exporta para CSV com delimitador ';' e codificação utf-8-sig para compatibilidade com Excel."""
        if not dados:
            return caminho_destino

        colunas = list(dados[0].keys())
        with open(caminho_destino, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=colunas, delimiter=";")
            writer.writeheader()
            writer.writerows(dados)

        return caminho_destino
