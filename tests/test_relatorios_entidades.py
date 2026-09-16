"""
Suite de Testes Automatizados para Módulos de Relatórios e Entidades do GeoAlvo.
Totalmente desacoplada de banco de dados real e janelas visuais (execução headless).
"""

import os
import sys
import unittest
import tempfile
from unittest.mock import MagicMock

# Assegura que a pasta raiz do GeoApoloV5 está no sys.path
PASTA_RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PASTA_RAIZ not in sys.path:
    sys.path.insert(0, PASTA_RAIZ)

from relatorios.models import TipoRelatorio, FiltroRelatorio, FormatoExportacao
from relatorios.generator import RelatorioGenerator, OPENPYXL_AVAILABLE
from relatorios.service import RelatorioService
from entidades.models import (
    EntidadeFiltro,
    ItemComparacao,
    DecisaoLinha,
    ResultadoOperacao,
    EntidadeEdicao,
)
from entidades.service import EntidadeService


class TestRelatorioGenerator(unittest.TestCase):
    """Testa geração de arquivos em Excel, CSV e HTML/PDF."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.dados_amostra = [
            {
                "Código": "001001",
                "Nome": "EMPRESA TESTE LTDA",
                "CPF/CNPJ": "12.345.678/0001-90",
                "Endereço": "RUA DAS FLORES",
                "Nº": "123",
                "Bairro": "CENTRO",
                "Cidade": "SÃO PAULO",
                "UF": "SP",
                "CEP": "01001-000",
                "Sincronizado Alvo": "Sim",
                "Data Cadastro": "14/09/2026",
            },
            {
                "Código": "001002",
                "Nome": "JOAO DA SILVA SAURO",
                "CPF/CNPJ": "123.456.789-00",
                "Endereço": "AVENIDA PAULISTA",
                "Nº": "1000",
                "Bairro": "BELA VISTA",
                "Cidade": "SÃO PAULO",
                "UF": "SP",
                "CEP": "01310-100",
                "Sincronizado Alvo": "Não",
                "Data Cadastro": "14/09/2026",
            },
        ]

    def tearDown(self):
        # Limpa arquivos temporários gerados
        for f in os.listdir(self.temp_dir):
            try:
                os.remove(os.path.join(self.temp_dir, f))
            except OSError:
                pass
        try:
            os.rmdir(self.temp_dir)
        except OSError:
            pass

    def test_gerar_excel_com_dados(self):
        """Verifica se a planilha Excel é gerada corretamente com cabeçalhos e linhas."""
        caminho_xlsx = os.path.join(self.temp_dir, "teste_relatorio.xlsx")
        resultado = RelatorioGenerator.gerar_excel(
            self.dados_amostra, "Relatório Geral de Entidades", caminho_xlsx
        )

        self.assertTrue(os.path.exists(resultado), "Arquivo Excel deve existir após geração.")
        self.assertGreater(os.path.getsize(resultado), 1000, "Arquivo Excel não deve estar vazio.")

        if OPENPYXL_AVAILABLE:
            import openpyxl
            wb = openpyxl.load_workbook(resultado)
            ws = wb.active
            self.assertIn("GEOALVO - RELATÓRIO GERAL DE ENTIDADES", str(ws["A1"].value))
            # Linha 4 tem os cabeçalhos
            cabecalhos = [cell.value for cell in ws[4]]
            self.assertIn("Código", cabecalhos)
            self.assertIn("Nome", cabecalhos)
            self.assertIn("CPF/CNPJ", cabecalhos)
            # Linha 5 tem os dados do primeiro registro
            valores_l5 = [str(cell.value) for cell in ws[5]]
            self.assertIn("001001", valores_l5)
            self.assertIn("EMPRESA TESTE LTDA", valores_l5)

    def test_gerar_excel_vazio(self):
        """Verifica se planilha com dados vazios é criada com aviso adequado."""
        caminho_xlsx = os.path.join(self.temp_dir, "teste_vazio.xlsx")
        resultado = RelatorioGenerator.gerar_excel([], "Relatório Vazio", caminho_xlsx)
        self.assertTrue(os.path.exists(resultado))

    def test_gerar_csv(self):
        """Verifica exportação em CSV com delimitador ';' e formatação correta."""
        caminho_csv = os.path.join(self.temp_dir, "teste_relatorio.csv")
        resultado = RelatorioGenerator.gerar_csv(self.dados_amostra, caminho_csv)

        self.assertTrue(os.path.exists(resultado), "Arquivo CSV deve ser criado.")
        with open(resultado, "r", encoding="utf-8-sig") as f:
            conteudo = f.read()
            self.assertIn("Código;Nome;CPF/CNPJ;Endereço", conteudo)
            self.assertIn("001001;EMPRESA TESTE LTDA", conteudo)
            self.assertIn("001002;JOAO DA SILVA SAURO", conteudo)

    def test_gerar_html_imprimivel(self):
        """Verifica a geração do documento HTML imprimível."""
        resultado = RelatorioGenerator.gerar_html_imprimivel(
            self.dados_amostra, "Relatório de Teste HTML", abrir_navegador=False
        )
        self.assertTrue(os.path.exists(resultado), "Arquivo HTML deve ser criado.")
        with open(resultado, "r", encoding="utf-8") as f:
            html = f.read()
            self.assertIn("<!DOCTYPE html>", html)
            self.assertIn("GEOALVO - RELATÓRIO DE TESTE HTML", html)
            self.assertIn("EMPRESA TESTE LTDA", html)
            self.assertIn("JOAO DA SILVA SAURO", html)
            self.assertIn("window.print()", html)
        try:
            os.remove(resultado)
        except OSError:
            pass


class TestEntidadeService(unittest.TestCase):
    """Testa regras de negócio e validações do EntidadeService."""

    def setUp(self):
        self.mock_repo = MagicMock()
        self.service = EntidadeService(self.mock_repo)

    def test_validar_dados_obrigatorios_sucesso(self):
        """Dados completos devem passar na validação."""
        dados = {
            "geoentcod": "001",
            "geoentnome": "PARÓQUIA SANTO ANTÔNIO",
            "geoentender": "PRAÇA DA MATRIZ",
            "geoenderno": "10",
            "geoentcep": "12345-000",
            "geocidcod": "001",
        }
        valido, msg = self.service.validar_dados(dados)
        self.assertTrue(valido)
        self.assertEqual(msg, "")

    def test_validar_dados_obrigatorios_falhas(self):
        """Campos ausentes devem acionar erro específico."""
        # Sem código
        valido, msg = self.service.validar_dados({"geoentnome": "TESTE"})
        self.assertFalse(valido)
        self.assertIn("código", msg.lower())

        # Sem nome
        valido, msg = self.service.validar_dados({"geoentcod": "001"})
        self.assertFalse(valido)
        self.assertIn("nome", msg.lower())

        # Sem endereço
        valido, msg = self.service.validar_dados({"geoentcod": "001", "geoentnome": "TESTE"})
        self.assertFalse(valido)
        self.assertIn("endereço", msg.lower())

    def test_pode_exportar_para_alvo(self):
        """Regras de exportação para API Alvo."""
        # Se estiver na base Alvo, exportação proibida
        pode, msg = self.service.pode_exportar_para_alvo("Alvo", None)
        self.assertFalse(pode)
        self.assertIn("já está na base alvo", msg.lower())

        # Se houver observações pendentes
        pode, msg = self.service.pode_exportar_para_alvo("GeoApolo", "Aguardando aprovação de doc")
        self.assertFalse(pode)
        self.assertIn("observações", msg.lower())

        # Válido para exportar
        pode, msg = self.service.pode_exportar_para_alvo("GeoApolo", "")
        self.assertTrue(pode)
        self.assertEqual(msg, "")

    def test_comparar_cadastros_identicos(self):
        """Cadastros iguais não devem gerar divergências."""
        sve = {
            "geoentnome": "ASSOCIAÇÃO BENEFICENTE",
            "Documento": "11.222.333/0001-44",
            "EntRgIe": "ISENTO",
            "tipolograd": "RUA",
            "geoentender": "CENTRAL",
            "geoenderno": "50",
            "geoentendercomp": "",
            "geoentbair": "CENTRO",
            "geoentcep": "11000-000",
            "cidnomecomp": "SANTOS",
            "ufsigla": "SP",
            "Email": "contato@associacao.org",
            "Telefone": "13 3333-4444",
            "geoentdataanivfund": "01/01/2000",
        }
        alvo = {
            "entnome": "ASSOCIAÇÃO BENEFICENTE",
            "EntCpfCgc": "11.222.333/0001-44",
            "EntRgIe": "ISENTO",
            "EntLograd": "RUA",
            "entender": "CENTRAL",
            "entenderno": "50",
            "EntEnderComp": "",
            "entbair": "CENTRO",
            "entcep": "11000-000",
            "cidnomecomp": "SANTOS",
            "ufsigla": "SP",
            "Email": "contato@associacao.org",
            "Telefone": "13 3333-4444",
            "EntDataAnivFund": "01/01/2000",
        }
        difs = self.service.comparar_cadastros(sve, alvo)
        self.assertEqual(len(difs), 0, "Cadastros idênticos devem retornar 0 divergências.")

    def test_comparar_cadastros_com_divergencias(self):
        """Cadastros divergentes devem identificar os campos alterados."""
        sve = {
            "geoentnome": "NOME NOVO NO SVE",
            "Documento": "11.222.333/0001-44",
            "geoentender": "RUA NOVA",
            "geoenderno": "999",
        }
        alvo = {
            "entnome": "NOME ANTIGO NO ALVO",
            "EntCpfCgc": "11.222.333/0001-44",
            "entender": "RUA ANTIGA",
            "entenderno": "999",
        }
        difs = self.service.comparar_cadastros(sve, alvo)
        # Nome e Endereço devem divergir
        rotulos_difs = [d.rotulo for d in difs]
        self.assertIn("Nome", rotulos_difs)
        self.assertIn("Endereço", rotulos_difs)
        self.assertNotIn("CPF / CNPJ", rotulos_difs)
        self.assertNotIn("Número", rotulos_difs)

    def test_gerar_payload_sobreposicao(self):
        """Verifica se o payload para API Alvo reflete as decisões do usuário."""
        itens = [
            ItemComparacao(
                indice_mapa=0,
                rotulo="Nome",
                campo_sve="geoentnome",
                campo_alvo="entnome",
                valor_sve="NOME DO SVE",
                valor_alvo="NOME DO ALVO",
                decisao=DecisaoLinha.MANTER_SVE,
            ),
            ItemComparacao(
                indice_mapa=4,
                rotulo="Endereço",
                campo_sve="geoentender",
                campo_alvo="entender",
                valor_sve="ENDERECO DO SVE",
                valor_alvo="ENDERECO DO ALVO",
                decisao=DecisaoLinha.MANTER_ALVO,
            ),
        ]
        payload = self.service.gerar_payload_sobreposicao(itens)
        self.assertEqual(payload["Operacao"], "A")
        self.assertEqual(payload["Entidade"]["Nome"], "NOME DO SVE")
        self.assertEqual(payload["Entidade"]["Endereco"], "ENDERECO DO ALVO")


class TestRelatorioService(unittest.TestCase):
    """Testa a geração de queries SQL do RelatorioService."""

    def test_queries_relatorio(self):
        mock_cursor = MagicMock()
        mock_cursor.description = [("Código",), ("Nome",)]
        mock_cursor.fetchall.return_value = [("001", "ENTIDADE 1"), ("002", "ENTIDADE 2")]

        mock_conn = MagicMock()
        mock_conn.cursor.return_value = mock_cursor

        service = RelatorioService(connection=mock_conn)

        # 1. ENTIDADES_GERAL
        filtro_geral = FiltroRelatorio(tipo_relatorio=TipoRelatorio.ENTIDADES_GERAL)
        dados = service.obter_dados_relatorio(filtro_geral)
        self.assertEqual(len(dados), 2)
        mock_cursor.execute.assert_called()
        sql_executado = mock_cursor.execute.call_args[0][0]
        self.assertIn("WITH (NOLOCK)", sql_executado)
        self.assertIn("USER_geoapolo_entidade", sql_executado)

        # 2. ENTIDADES_SINCRONIZADAS
        filtro_sinc = FiltroRelatorio(tipo_relatorio=TipoRelatorio.ENTIDADES_SINCRONIZADAS)
        dados_sinc = service.obter_dados_relatorio(filtro_sinc)
        self.assertEqual(len(dados_sinc), 2)
        sql_sinc = mock_cursor.execute.call_args[0][0]
        self.assertIn("atualizou_apolo = 'S'", sql_sinc)

        # 3. ENTIDADES_PENDENTES
        filtro_pend = FiltroRelatorio(tipo_relatorio=TipoRelatorio.ENTIDADES_PENDENTES)
        dados_pend = service.obter_dados_relatorio(filtro_pend)
        sql_pend = mock_cursor.execute.call_args[0][0]
        self.assertIn("atualizou_apolo = 'N'", sql_pend)

        # 4. OCORRENCIAS_SISTEMA
        filtro_ocor = FiltroRelatorio(tipo_relatorio=TipoRelatorio.OCORRENCIAS_SISTEMA)
        dados_ocor = service.obter_dados_relatorio(filtro_ocor)
        sql_ocor = mock_cursor.execute.call_args[0][0]
        self.assertIn("ocorcod IS NOT NULL", sql_ocor)


if __name__ == "__main__":
    unittest.main()
